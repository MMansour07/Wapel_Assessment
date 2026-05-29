#!/usr/bin/env python3
"""Convert Markdown files to .docx with basic structure preservation."""

from __future__ import annotations

import argparse
from pathlib import Path

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt


def add_runs(paragraph, node: Tag | NavigableString) -> None:
    if isinstance(node, NavigableString):
        text = str(node)
        if text:
            paragraph.add_run(text)
        return

    if not isinstance(node, Tag):
        return

    name = node.name.lower()
    if name == "br":
        paragraph.add_run("\n")
        return

    if name in {"strong", "b"}:
        run = paragraph.add_run(node.get_text())
        run.bold = True
        return

    if name in {"em", "i"}:
        run = paragraph.add_run(node.get_text())
        run.italic = True
        return

    if name == "code":
        run = paragraph.add_run(node.get_text())
        run.font.name = "Courier New"
        run.font.size = Pt(10)
        return

    for child in node.children:
        add_runs(paragraph, child)


def add_paragraph_from_node(doc: Document, node: Tag, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    for child in node.children:
        add_runs(paragraph, child)


def convert_element(doc: Document, element: Tag) -> None:
    name = element.name.lower()

    if name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
        level = int(name[1])
        add_paragraph_from_node(doc, element, style=f"Heading {min(level, 9)}")
        return

    if name == "p":
        add_paragraph_from_node(doc, element)
        return

    if name == "blockquote":
        add_paragraph_from_node(doc, element, style="Intense Quote")
        return

    if name in {"ul", "ol"}:
        list_style = "List Bullet" if name == "ul" else "List Number"
        for li in element.find_all("li", recursive=False):
            add_paragraph_from_node(doc, li, style=list_style)
        return

    if name == "pre":
        code_text = element.get_text()
        paragraph = doc.add_paragraph(style="No Spacing")
        run = paragraph.add_run(code_text.rstrip("\n"))
        run.font.name = "Courier New"
        run.font.size = Pt(10)
        return

    if name == "hr":
        doc.add_paragraph("-" * 64)
        return

    if name == "table":
        rows = element.find_all("tr")
        if not rows:
            return

        max_cols = 0
        row_cells = []
        for row in rows:
            cells = row.find_all(["th", "td"], recursive=False)
            row_cells.append(cells)
            if len(cells) > max_cols:
                max_cols = len(cells)

        table = doc.add_table(rows=len(rows), cols=max_cols)
        table.style = "Table Grid"

        for r_idx, cells in enumerate(row_cells):
            for c_idx, cell in enumerate(cells):
                table.cell(r_idx, c_idx).text = cell.get_text(separator=" ", strip=True)

        doc.add_paragraph()
        return

    # Fallback for unhandled tags.
    if element.get_text(strip=True):
        add_paragraph_from_node(doc, element)


def convert_markdown_to_docx(input_path: Path, output_path: Path) -> None:
    markdown_text = input_path.read_text(encoding="utf-8")
    html = markdown.markdown(
        markdown_text,
        extensions=["tables", "fenced_code", "sane_lists"],
    )
    soup = BeautifulSoup(html, "html.parser")

    doc = Document()

    for child in soup.contents:
        if isinstance(child, Tag):
            convert_element(doc, child)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Convert Markdown to DOCX.")
    parser.add_argument("input_markdown", type=Path, help="Source markdown file path.")
    parser.add_argument("output_docx", type=Path, help="Output docx file path.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    convert_markdown_to_docx(args.input_markdown, args.output_docx)


if __name__ == "__main__":
    main()
